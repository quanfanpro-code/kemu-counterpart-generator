"""验证批处理不会漏文件、覆盖原账或因单份坏账中断。"""
import csv
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook


def ledger(path, renamed=False, invalid=False):
    df = pd.DataFrame({
        "会计月": [1, 1], "凭证编号": ["001", "001"],
        "一级科目": ["银行存款", "应收账款"],
        "借方发生额": ["错误" if invalid else 100, 0],
        "贷方发生额": [0, 100], "记账时间": ["2026-10-01 23:00", "2026-10-01 23:00"],
    })
    if renamed:
        df.rename(columns={"会计月": "期间"}, inplace=True)
    df.to_excel(path, index=False)


def test_根层递归与排除历史结果(tmp_path):
    from src.pipeline.批处理 import scan_ledger_files
    sub = tmp_path / "子目录"
    sub.mkdir()
    old = tmp_path / "批处理结果_20261001_120000"
    old.mkdir()
    for path in [tmp_path / "账.xlsx", sub / "子账.XLS", tmp_path / "~$账.xlsx",
                 tmp_path / "账生成对方科目.xlsx", old / "历史账.xlsx"]:
        path.touch()
    assert [p.name for p in scan_ledger_files(tmp_path)] == ["账.xlsx"]
    assert [p.relative_to(tmp_path) for p in scan_ledger_files(tmp_path, True)] == [
        Path("子目录") / "子账.XLS", Path("账.xlsx")]


def test_坏账继续输出分层清单且输入不变(tmp_path):
    from src.pipeline.批处理 import run_batch
    (tmp_path / "子目录").mkdir()
    (tmp_path / "00损坏.xlsx").write_bytes(b"broken")
    ledger(tmp_path / "01非法金额.xlsx", invalid=True)
    ledger(tmp_path / "账.xlsx")
    ledger(tmp_path / "子目录" / "账.xlsx")
    originals = {p: p.read_bytes() for p in tmp_path.rglob("*.xlsx")}
    result = run_batch(tmp_path, recursive=True)
    assert [r["状态"] for r in result["records"]] == ["失败", "失败", "成功", "成功"]
    assert "金额" in result["records"][1]["原因"]
    for relative in [Path("账生成对方科目.xlsx"), Path("子目录") / "账生成对方科目.xlsx"]:
        with pd.ExcelFile(result["output_dir"] / relative) as book:
            df = pd.read_excel(book, sheet_name="生成结果").dropna(axis=1, how="all")
        assert df["对方科目"].tolist() == ["应收账款", "银行存款"]
    assert all(p.read_bytes() == data for p, data in originals.items())
    with (result["output_dir"] / "批处理清单.csv").open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    assert [r["状态"] for r in rows] == ["失败", "失败", "成功", "成功"]
    second = run_batch(tmp_path, recursive=True)
    assert second["output_dir"] != result["output_dir"]
    assert len(second["records"]) == 4


def test_同格式配置与筛选整批复用(tmp_path):
    from src.pipeline.批处理 import run_batch
    for name in ["甲.xlsx", "乙.xlsx"]:
        ledger(tmp_path / name, renamed=True)
    requests, screens = [], []
    def mapping(columns, required):
        requests.append(columns)
        return {col: "期间" if col == "会计月" else
                "无/不适用" if col == "凭证种类" else col for col in required}
    def screening(df):
        screens.append(df.attrs["source_path"])
        return {"time": {"column": "记账时间", "night": True, "start": "08:00", "end": "18:00"}}
    result = run_batch(tmp_path, mapping_dialog=mapping, screening_dialog=screening)
    assert len(requests) == len(screens) == 1
    assert all(r["状态"] == "成功" for r in result["records"])
    for path in result["output_dir"].glob("*.xlsx"):
        book = load_workbook(path, read_only=True)
        assert "入账时间筛选" in book.sheetnames
        book.close()


def test_取消一种格式其他格式继续(tmp_path):
    from src.pipeline.批处理 import run_batch
    ledger(tmp_path / "01.xlsx", renamed=True)
    ledger(tmp_path / "02.xlsx", renamed=True)
    ledger(tmp_path / "03.xlsx")
    requests = []
    def cancel(columns, required):
        requests.append(columns)
        return None
    result = run_batch(tmp_path, mapping_dialog=cancel)
    assert [r["状态"] for r in result["records"]] == ["失败", "失败", "成功"]
    assert len(requests) == 1


def test_空目录不创建结果(tmp_path):
    from src.pipeline.批处理 import run_batch
    result = run_batch(tmp_path)
    assert result["output_dir"] is None and result["records"] == []
    assert list(tmp_path.iterdir()) == []


def test_筛选取消记录未处理(tmp_path):
    from src.pipeline.批处理 import run_batch
    ledger(tmp_path / "甲.xlsx")
    ledger(tmp_path / "乙.xlsx")
    result = run_batch(tmp_path, screening_dialog=lambda df: None)
    assert [r["状态"] for r in result["records"]] == ["未处理", "未处理"]
    assert not list(result["output_dir"].glob("*.xlsx"))


def test_非目录明确报错(tmp_path):
    from src.pipeline.批处理 import run_batch
    with pytest.raises(ValueError, match="文件夹"):
        run_batch(tmp_path / "不存在")


def test_同名不同格式与后缀巧合同名不覆盖(tmp_path):
    from src.pipeline.批处理 import run_batch
    ledger(tmp_path / "账.xlsx")
    (tmp_path / "账.xls").write_bytes((tmp_path / "账.xlsx").read_bytes())
    ledger(tmp_path / "账_xls.xlsx")
    result = run_batch(tmp_path)
    assert all(row["状态"] == "成功" for row in result["records"])
    assert len({row["结果文件"] for row in result["records"]}) == 3
    assert len(list(result["output_dir"].glob("*.xlsx"))) == 3


def test_保存失败不发布正式结果且继续(tmp_path, monkeypatch):
    from src.pipeline import 批处理 as batch
    from src.io import writer
    ledger(tmp_path / "甲.xlsx")
    ledger(tmp_path / "乙.xlsx")
    original = writer.save_output_file
    def save(path, *args, **kwargs):
        if "乙" in Path(path).name:
            batch.logger.error("文件被占用，无法保存")
            return False
        return original(path, *args, **kwargs)
    monkeypatch.setattr(writer, "save_output_file", save)
    result = batch.run_batch(tmp_path)
    statuses = {row["来源文件"]: row["状态"] for row in result["records"]}
    assert statuses == {"乙.xlsx": "失败", "甲.xlsx": "成功"}
    assert not (result["output_dir"] / "乙生成对方科目.xlsx").exists()
    assert (result["output_dir"] / "甲生成对方科目.xlsx").exists()


def test_清单防公式且仅处理首张工作表(tmp_path):
    from src.pipeline.批处理 import run_batch
    source = tmp_path / "=账.xlsx"
    ledger(source)
    with pd.ExcelWriter(source, engine="openpyxl", mode="a") as book:
        pd.DataFrame({"非序时账": [1]}).to_excel(book, sheet_name="其他", index=False)
    result = run_batch(tmp_path)
    assert result["records"][0]["状态"] == "成功"
    with (result["output_dir"] / "批处理清单.csv").open(encoding="utf-8-sig", newline="") as f:
        row = next(csv.DictReader(f))
    assert row["来源文件"] == "'=账.xlsx"


def test_不同格式分别确认(tmp_path):
    from src.pipeline.批处理 import run_batch
    ledger(tmp_path / "甲.xlsx", renamed=True)
    ledger(tmp_path / "乙.xlsx", renamed=True)
    frame = pd.read_excel(tmp_path / "乙.xlsx").rename(columns={"期间": "月份"})
    frame.to_excel(tmp_path / "乙.xlsx", index=False)
    requests = []
    def mapping(columns, required):
        requests.append(columns)
        return {name: ("期间" if "期间" in columns else "月份") if name == "会计月"
                else "无/不适用" if name == "凭证种类" else name for name in required}
    result = run_batch(tmp_path, mapping_dialog=mapping)
    assert len(requests) == 2
    assert all(row["状态"] == "成功" for row in result["records"])


def test_凭证组出错保存结果但不误报全部成功(tmp_path, monkeypatch):
    from src.pipeline import orchestrator
    from src.pipeline.批处理 import run_batch
    ledger(tmp_path / "账.xlsx")
    original = orchestrator.process_group
    def fail_group(task):
        if str(task[0][-1]) == "2":
            raise ValueError("测试凭证组处理失败")
        return original(task)
    frame = pd.read_excel(tmp_path / "账.xlsx")
    other = frame.copy()
    other["凭证编号"] = "002"
    pd.concat([frame, other]).to_excel(tmp_path / "账.xlsx", index=False)
    monkeypatch.setattr(orchestrator, "process_group", fail_group)
    result = run_batch(tmp_path)
    assert result["records"][0]["状态"] == "部分失败"
    book = load_workbook(result["output_dir"] / result["records"][0]["结果文件"], read_only=True)
    assert "失败分组" in book.sheetnames
    book.close()
