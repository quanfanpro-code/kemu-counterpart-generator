import tkinter as tk
from src.gui.筛选设置 import ScreeningDialog
from tests.test_凭证筛选 import frame


def test_桌面设置选择科目方向并读取配置():
    root=tk.Tk()
    root.withdraw()
    try:
        dialog=ScreeningDialog(root,frame(),True,True)
        dialog.vars["date"].set("记账时间")
        dialog.vars["direction"].set("贷方")
        dialog.vars["limit"].set("100000")
        dialog.vars["minimum"].set("90000")
        dialog.subjects.selection_set(0)
        dialog.vars["split_date"].set("记账时间")
        dialog.vars["party"].set("对方单位")
        dialog.update()
        assert dialog.options()["split"]["subjects"]==["银行存款"]
        dialog.accept()
        assert dialog.result["split"]["direction"]=="贷方"
        assert dialog.result["time"]["holiday"] is True
    finally:
        root.destroy()



def test_从桌面主窗口勾选到实际导出(tmp_path, monkeypatch):
    import sys
    import time
    from src.gui import app
    from src.gui import 筛选设置 as settings
    from openpyxl import load_workbook
    df=frame()
    df["借方发生额"]=df["贷方发生额"]
    source=tmp_path/"桌面入口.xlsx"
    df.to_excel(source,index=False)
    target=tmp_path/"桌面入口生成对方科目.xlsx"
    monkeypatch.setattr(app.filedialog,"askopenfilename",lambda **_:str(source))
    original_button=app._make_button
    original_init=settings.ScreeningDialog.__init__
    errors=[]
    def dialog_init(self,*args,**kwargs):
        original_init(self,*args,**kwargs)
        def fill():
            self.vars["date"].set("记账时间")
            self.accept()
        self.after(50,fill)
    monkeypatch.setattr(settings.ScreeningDialog,"__init__",dialog_init)
    def descendants(widget):
        for child in widget.winfo_children():
            yield child
            yield from descendants(child)
    def make_button(master,**kwargs):
        button=original_button(master,**kwargs)
        if kwargs.get("text")=="📁 选择Excel文件":
            root=master.winfo_toplevel()
            deadline=time.monotonic()+25
            def poll():
                if target.exists() and target.stat().st_size>0:
                    try:
                        book=load_workbook(target)
                        assert "入账时间筛选" in book.sheetnames
                        book.close()
                    except (OSError,ValueError):
                        root.after(100,poll);return
                    except Exception as exc:
                        errors.append(str(exc))
                    root.destroy()
                elif time.monotonic()>deadline:
                    errors.append("桌面处理超时");root.destroy()
                else:
                    root.after(100,poll)
            def choose():
                try:
                    checks=[w for w in descendants(root) if isinstance(w,app.ctk.CTkCheckBox)]
                    assert len(checks)==2 and all(w.get()==0 for w in checks)
                    checks[0].select()
                    root.after(100,poll)
                    button.invoke()
                except Exception as exc:
                    errors.append(str(exc));root.destroy()
            root.after(100,choose)
        return button
    monkeypatch.setattr(app,"_make_button",make_button)
    stdout,stderr=sys.stdout,sys.stderr
    try:
        app.run_gui()
    finally:
        sys.stdout,sys.stderr=stdout,stderr
    assert not errors,errors
    assert target.exists()


def test_设置不完整留在窗口且可选择只做原有分析():
    root=tk.Tk()
    root.withdraw()
    try:
        dialog=ScreeningDialog(root,frame(),False,True)
        dialog.accept()
        assert dialog.winfo_exists()
        assert dialog.result is None
        dialog.skip()
        assert dialog.result=={}
    finally:
        root.destroy()
