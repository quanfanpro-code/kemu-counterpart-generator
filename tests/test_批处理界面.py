"""通过真实桌面入口验证递归开关、运行保护和批次导出。"""
import sys
import time

import pytest
from tests.test_文件夹批处理 import ledger


@pytest.mark.parametrize("recursive", [False, True])
def test_文件夹入口递归开关与运行保护(tmp_path, monkeypatch, recursive):
    from src.gui import app
    ledger(tmp_path / "根账.xlsx")
    (tmp_path / "子目录").mkdir()
    ledger(tmp_path / "子目录" / "子账.xlsx")
    monkeypatch.setattr(app.filedialog, "askdirectory", lambda **_: str(tmp_path))
    original_button = app._make_button
    original_mainloop = app.ctk.CTk.mainloop
    buttons, errors = {}, []
    def make_button(master, **kwargs):
        button = original_button(master, **kwargs)
        buttons[kwargs["text"]] = button
        return button
    def descendants(widget):
        for child in widget.winfo_children():
            yield child
            yield from descendants(child)
    def mainloop(root, *args, **kwargs):
        deadline = time.monotonic() + 20
        def poll():
            finished = any(str(w.cget("text")).startswith("批处理完成")
                           for w in descendants(root) if isinstance(w, app.ctk.CTkLabel))
            if finished:
                root.destroy()
            elif time.monotonic() > deadline:
                errors.append("批处理界面超时")
                root.destroy()
            else:
                root.after(50, poll)
        def choose():
            try:
                checks = [w for w in descendants(root) if isinstance(w, app.ctk.CTkCheckBox)
                          and w.cget("text") == "包含子文件夹"]
                assert len(checks) == 1 and checks[0].get() == 0
                if recursive:
                    checks[0].select()
                assert "选择文件夹" in buttons
                buttons["选择文件夹"].invoke()
                assert buttons["选择文件夹"].cget("state") == "disabled"
                root.after(50, poll)
            except Exception as exc:
                errors.append(repr(exc))
                root.destroy()
        root.after(100, choose)
        original_mainloop(root, *args, **kwargs)
    monkeypatch.setattr(app, "_make_button", make_button)
    monkeypatch.setattr(app.ctk.CTk, "mainloop", mainloop)
    stdout, stderr = sys.stdout, sys.stderr
    try:
        app.run_gui()
    finally:
        sys.stdout, sys.stderr = stdout, stderr
    assert not errors, errors
    folders = list(tmp_path.glob("批处理结果_*"))
    assert len(folders) == 1
    assert (folders[0] / "根账生成对方科目.xlsx").exists()
    assert (folders[0] / "子目录" / "子账生成对方科目.xlsx").exists() is recursive


def test_批次弹窗主线程配置一次并实际筛选(tmp_path, monkeypatch):
    import threading
    from src.gui import app, 批处理设置, 筛选设置
    from openpyxl import load_workbook
    for name in ("甲.xlsx", "乙.xlsx"):
        ledger(tmp_path / name, renamed=True)
    monkeypatch.setattr(app.filedialog, "askdirectory", lambda **_: str(tmp_path))
    main_thread = threading.get_ident()
    mapping_requests, screening_requests, errors = [], [], []
    original_mapping = 批处理设置.BatchMappingDialog.__init__
    original_screen = 筛选设置.ScreeningDialog.__init__
    original_loop = app.ctk.CTk.mainloop
    def configure_mapping(self, *args):
        original_mapping(self, *args)
        mapping_requests.append(threading.get_ident())
        def fill():
            self.choices["会计月"].set("期间")
            self.accept()
        self.after(30, fill)
    def configure_screen(self, *args):
        original_screen(self, *args)
        screening_requests.append(threading.get_ident())
        def fill():
            self.vars["date"].set("记账时间")
            self.accept()
        self.after(30, fill)
    def descendants(widget):
        for child in widget.winfo_children():
            yield child
            yield from descendants(child)
    def loop(root, *args, **kwargs):
        deadline = time.monotonic() + 20
        def poll():
            if any(isinstance(w, app.ctk.CTkLabel) and str(w.cget("text")).startswith("批处理完成")
                   for w in descendants(root)):
                root.destroy()
            elif time.monotonic() > deadline:
                errors.append("配置后批次超时")
                root.destroy()
            else:
                root.after(50, poll)
        def choose():
            checks = [w for w in descendants(root) if isinstance(w, app.ctk.CTkCheckBox)
                      and str(w.cget("text")).startswith("按入账时间")]
            checks[0].select()
            buttons = [w for w in descendants(root) if isinstance(w, app.ctk.CTkButton)
                       and w.cget("text") == "选择文件夹"]
            buttons[0].invoke()
            root.after(50, poll)
        root.after(100, choose)
        original_loop(root, *args, **kwargs)
    monkeypatch.setattr(批处理设置.BatchMappingDialog, "__init__", configure_mapping)
    monkeypatch.setattr(筛选设置.ScreeningDialog, "__init__", configure_screen)
    monkeypatch.setattr(app.ctk.CTk, "mainloop", loop)
    stdout, stderr = sys.stdout, sys.stderr
    try:
        app.run_gui()
    finally:
        sys.stdout, sys.stderr = stdout, stderr
    assert not errors, errors
    assert mapping_requests == screening_requests == [main_thread]
    outputs = list(tmp_path.glob("批处理结果_*\\*.xlsx"))
    assert len(outputs) == 2
    for path in outputs:
        book = load_workbook(path, read_only=True)
        assert "入账时间筛选" in book.sheetnames
        book.close()


def test_列对应不完整或重复不接受():
    import tkinter as tk
    from src.gui.批处理设置 import BatchMappingDialog
    root = tk.Tk()
    root.withdraw()
    try:
        dialog = BatchMappingDialog(root, ["甲", "乙"], ["会计月", "凭证编号"])
        dialog.accept()
        assert dialog.result is None and dialog.winfo_exists()
        dialog.choices["会计月"].set("甲")
        dialog.choices["凭证编号"].set("甲")
        dialog.accept()
        assert dialog.result is None and dialog.winfo_exists()
        dialog.choices["凭证编号"].set("乙")
        dialog.accept()
        assert dialog.result == {"会计月": "甲", "凭证编号": "乙"}
    finally:
        root.destroy()
