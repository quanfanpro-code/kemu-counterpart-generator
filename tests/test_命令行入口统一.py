# -*- coding: utf-8 -*-
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
MAIN = ROOT / "main.py"
CLI = ROOT / "cli" / "序时账分析器命令行版.py"


def run_entry(entry, *args):
    return subprocess.run(
        [sys.executable, "-B", "-X", "utf8", str(entry), *map(str, args)],
        cwd=ROOT, capture_output=True, text=True, encoding="utf-8",
        errors="replace", timeout=90,
    )


@pytest.mark.parametrize("entry", [MAIN, CLI])
def test_无头入口显示当前版本(entry):
    finished = run_entry(entry, "--help")
    assert finished.returncode == 0
    assert "v2.5.0" in finished.stdout


def test_主程序无头模式支持简版输出(tmp_path):
    from openpyxl import load_workbook
    from tests.test_文件夹批处理 import ledger

    source = tmp_path / "账.xlsx"
    output = tmp_path / "简版.xlsx"
    ledger(source)
    finished = run_entry(
        MAIN, source, output, "--no-gui", "--mode", "summary", "--log-level", "ERROR",
    )
    assert finished.returncode == 0, finished.stdout + finished.stderr
    book = load_workbook(output)
    assert book.sheetnames == ["生成结果"]
    assert not book["生成结果"]._charts
    book.close()


def test_共享流水线拒绝未知输出模式(tmp_path):
    from src.pipeline.orchestrator import run_processing_pipeline

    with pytest.raises(ValueError, match="output_mode"):
        run_processing_pipeline(None, 10000, str(tmp_path / "bad.xlsx"), output_mode="bad")


def test_旧CLI完整输出与主程序一致(tmp_path):
    from openpyxl import load_workbook
    from tests.test_文件夹批处理 import ledger

    source = tmp_path / "账.xlsx"
    output = tmp_path / "完整.xlsx"
    ledger(source)
    finished = run_entry(CLI, source, output, "--mode", "full", "--log-level", "ERROR")
    assert finished.returncode == 0, finished.stdout + finished.stderr
    book = load_workbook(output)
    assert "班福分析" in book.sheetnames
    assert len(book["班福分析"]._charts) == 1
    book.close()


def test_旧CLI无参数不启动图形界面():
    try:
        finished = subprocess.run(
            [sys.executable, "-B", "-X", "utf8", str(CLI)],
            cwd=ROOT, capture_output=True, text=True, encoding="utf-8",
            errors="replace", timeout=3,
        )
    except subprocess.TimeoutExpired:
        pytest.fail("旧 CLI 无参数时启动了图形界面")
    assert finished.returncode == 2
    assert "--no-gui 模式需要同时提供输入和输出文件路径" in finished.stderr
