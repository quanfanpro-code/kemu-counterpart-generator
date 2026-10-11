"""验证手动开始、安全停止和列映射可读性。"""
import csv
import sys
import threading
import time
import tkinter as tk
import tkinter.font as tkfont
from tkinter import ttk

import pytest
from openpyxl import load_workbook

from tests.test_文件夹批处理 import ledger


@pytest.mark.parametrize("pre_stopped", [False, True])
def test_批处理停止保留完整文件与未处理清单(tmp_path, pre_stopped):
    from src.pipeline.批处理 import run_batch
    for name in ("01.xlsx", "02.xlsx"):
        ledger(tmp_path / name)
    stop = threading.Event()
    if pre_stopped:
        stop.set()
    def progress(percent, message, phase=None):
        if phase == "文件输出":
            stop.set()
    result = run_batch(tmp_path, stop_event=stop, progress_callback=progress)
    expected = ["未处理", "未处理"] if pre_stopped else ["成功", "未处理"]
    assert [r["状态"] for r in result["records"]] == expected
    assert "停止" in result["records"][-1]["原因"]
    outputs = list(result["output_dir"].glob("*.xlsx"))
    assert len(outputs) == (0 if pre_stopped else 1)
    for path in outputs:
        book = load_workbook(path, read_only=True)
        assert "生成结果" in book.sheetnames
        book.close()
    with (result["output_dir"] / "批处理清单.csv").open(encoding="utf-8-sig", newline="") as stream:
        assert [r["状态"] for r in csv.DictReader(stream)] == expected


def descendants(widget):
    for child in widget.winfo_children():
        yield child
        yield from descendants(child)


@pytest.mark.parametrize("folder", [False, True])
def test_选择不运行点击开始安全停止并恢复(tmp_path, monkeypatch, folder):
    from src.gui import app
    from src.gui import 批处理设置
    from src.io import writer
    ledger(tmp_path / "01.xlsx")
    ledger(tmp_path / "02.xlsx")
    monkeypatch.setattr(app.filedialog, "askopenfilename", lambda **_: str(tmp_path / "01.xlsx"))
    monkeypatch.setattr(app.filedialog, "askdirectory", lambda **_: str(tmp_path))
    original_mapping = 批处理设置.BatchMappingDialog.__init__
    def mapping(self, *args, **kwargs):
        original_mapping(self, *args, **kwargs)
        self.after(30, self.accept)
    monkeypatch.setattr(批处理设置.BatchMappingDialog, "__init__", mapping)
    saving, release = threading.Event(), threading.Event()
    original_save = writer.save_output_file
    def save(*args, **kwargs):
        saving.set()
        assert release.wait(10), "等待停止按钮超时"
        return original_save(*args, **kwargs)
    monkeypatch.setattr(writer, "save_output_file", save)
    original_loop = app.ctk.CTk.mainloop
    errors = []
    def loop(root, *args, **kwargs):
        buttons = {w.cget("text"): w for w in descendants(root)
                   if isinstance(w, app.ctk.CTkButton)}
        deadline = time.monotonic() + 15
        def fail(exc):
            errors.append(repr(exc))
            release.set()
            root.destroy()
        def poll():
            try:
                assert time.monotonic() < deadline, "处理控制超时"
                if saving.is_set() and not release.is_set():
                    assert buttons["开始处理"].cget("state") == "disabled"
                    buttons["停止处理"].invoke()
                    assert buttons["停止处理"].cget("state") == "disabled"
                    release.set()
                if release.is_set() and buttons["开始处理"].cget("state") == "normal":
                    texts = [str(w.cget("text")) for w in descendants(root)
                             if isinstance(w, app.ctk.CTkLabel)]
                    assert any("已停止" in text for text in texts), texts
                    root.destroy()
                else:
                    root.after(30, poll)
            except Exception as exc:
                fail(exc)
        def choose():
            try:
                assert buttons["开始处理"].cget("state") == "disabled"
                assert buttons["停止处理"].cget("state") == "disabled"
                buttons["选择文件夹" if folder else "📁 选择Excel文件"].invoke()
                assert not saving.is_set()
                assert not list(tmp_path.glob("批处理结果_*"))
                assert not (tmp_path / "01生成对方科目.xlsx").exists()
                buttons["开始处理"].invoke()
                root.after(30, poll)
            except Exception as exc:
                fail(exc)
        root.after(100, choose)
        original_loop(root, *args, **kwargs)
    monkeypatch.setattr(app.ctk.CTk, "mainloop", loop)
    stdout, stderr = sys.stdout, sys.stderr
    try:
        app.run_gui()
    finally:
        release.set()
        sys.stdout, sys.stderr = stdout, stderr
    assert not errors, errors
    outputs = list(tmp_path.glob("批处理结果_*\\*.xlsx")) if folder else [
        tmp_path / "01生成对方科目.xlsx"]
    assert len(outputs) == 1
    book = load_workbook(outputs[0], read_only=True)
    assert "生成结果" in book.sheetnames
    book.close()


@pytest.mark.parametrize("scaling", [1.0, 1.25, 1.5])
def test_映射不同缩放下完整可见(scaling):
    import customtkinter as ctk
    from src.gui.批处理设置 import BatchMappingDialog
    ctk.set_widget_scaling(scaling)
    root = ctk.CTk()
    try:
        root.update()
        fields = ["会计月", "凭证编号", "一级科目", "借方发生额", "贷方发生额", "凭证种类"]
        dialog = BatchMappingDialog(root, fields, fields)
        root.update()
        assert len(dialog.choices) == 6
        buttons = [w for w in descendants(dialog) if isinstance(w, ttk.Button)]
        assert len(buttons) == 2
        for widget in buttons:
            assert widget.winfo_ismapped()
            assert widget.winfo_rooty() + widget.winfo_height() <= dialog.winfo_rooty() + dialog.winfo_height()
        assert dialog.mapping_font.metrics("linespace") >= 22
        dialog.destroy()
    finally:
        root.destroy()
        ctk.set_widget_scaling(1.0)


def test_当前文件失败且请求停止不误报成功(tmp_path, monkeypatch):
    from src.pipeline import 批处理 as batch
    from src.io import writer
    for name in ("01.xlsx", "02.xlsx"):
        ledger(tmp_path / name)
    stop = threading.Event()
    def fail(*args, **kwargs):
        stop.set()
        return False
    monkeypatch.setattr(writer, "save_output_file", fail)
    result = batch.run_batch(tmp_path, stop_event=stop)
    assert [r["状态"] for r in result["records"]] == ["失败", "未处理"]
    assert "保存失败" in result["records"][0]["原因"]
    assert "停止" in result["records"][1]["原因"]
    assert not list(result["output_dir"].glob("*.xlsx"))


def test_映射字体与下拉列表可读():
    from src.gui.批处理设置 import BatchMappingDialog
    root = tk.Tk()
    try:
        dialog = BatchMappingDialog(root, ["期间", "凭证号"], ["会计月", "凭证编号"])
        root.update()
        combos = [w for w in descendants(dialog) if isinstance(w, ttk.Combobox)]
        for combo in combos:
            font = tkfont.Font(root=root, font=combo.cget("font"))
            assert font.metrics("linespace") >= 22
            assert combo.winfo_height() >= font.metrics("linespace") + 4
            popdown = root.tk.call("ttk::combobox::PopdownWindow", str(combo))
            list_font = root.tk.call(str(popdown) + ".f.l", "cget", "-font")
            assert tkfont.Font(root=root, font=list_font).metrics("linespace") >= 22
        dialog.destroy()
    finally:
        root.destroy()


@pytest.mark.parametrize("renamed", [False, True])
def test_单文件映射始终可核对五个必需字段(tmp_path, monkeypatch, renamed):
    from src.gui import app, 批处理设置
    ledger(tmp_path / "账.xlsx", renamed=renamed)
    monkeypatch.setattr(app.filedialog, "askopenfilename", lambda **_: str(tmp_path / "账.xlsx"))
    original_mapping = 批处理设置.BatchMappingDialog.__init__
    original_loop = app.ctk.CTk.mainloop
    errors, shown = [], []
    def mapping(self, *args, **kwargs):
        original_mapping(self, *args, **kwargs)
        shown.append(set(self.choices))
        def accept():
            if renamed:
                self.choices["会计月"].set("期间")
            self.accept()
        self.after(30, accept)
    def loop(root, *args, **kwargs):
        def choose():
            try:
                buttons = {w.cget("text"): w for w in descendants(root)
                           if isinstance(w, app.ctk.CTkButton)}
                buttons["📁 选择Excel文件"].invoke()
                assert shown == [{"会计月", "凭证编号", "一级科目", "借方发生额",
                                  "贷方发生额", "凭证种类"}]
                assert buttons["开始处理"].cget("state") == "normal"
                assert not (tmp_path / "账生成对方科目.xlsx").exists()
            except Exception as exc:
                errors.append(repr(exc))
            finally:
                root.destroy()
        root.after(100, choose)
        original_loop(root, *args, **kwargs)
    monkeypatch.setattr(批处理设置.BatchMappingDialog, "__init__", mapping)
    monkeypatch.setattr(app.ctk.CTk, "mainloop", loop)
    stdout, stderr = sys.stdout, sys.stderr
    try:
        app.run_gui()
    finally:
        sys.stdout, sys.stderr = stdout, stderr
    assert not errors, errors
