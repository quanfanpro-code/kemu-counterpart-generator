"""从原始序时账筛选记录。默认关闭，不改变对方科目生成算法。"""
import re
from collections import defaultdict
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation

import pandas as pd

from src.日历 import load_calendar, day_type


def _datetime(value):
    if pd.isna(value):
        return None, False
    if isinstance(value, (datetime, pd.Timestamp)):
        return value.to_pydatetime() if isinstance(value, pd.Timestamp) else value, value.time()!=time()
    if isinstance(value, date):
        return datetime.combine(value,time()),False
    text=str(value).strip()
    # 只有完整年月日才解释为日期，不能将会计月或Excel序号猜成日期。
    if not re.match(r"^\d{4}[-/年]\d{1,2}[-/月]\d{1,2}(?:日|\b|T)",text):
        return None,False
    parsed=pd.to_datetime(text.replace("年","-").replace("月","-").replace("日",""),errors="coerce")
    if pd.isna(parsed) or parsed.tzinfo is not None:
        return None,False
    return parsed.to_pydatetime(),bool(re.search(r"(?:T|\s)\d{1,2}:\d{2}",text))


def _precision(df, column):
    path=df.attrs.get("source_path")
    if not path or not str(path).lower().endswith(".xlsx"):
        return {}
    from openpyxl import load_workbook
    book=load_workbook(path,read_only=True,data_only=True)
    try:
        sheet=book.worksheets[0]
        header=next(sheet.iter_rows(max_row=1,values_only=True))
        original=df.attrs.get("column_mapping",{}).get(column,column)
        if original not in header:
            return {}
        col=header.index(original)+1
        result={}
        for row in sheet.iter_rows(min_row=2,min_col=col,max_col=col):
            cell=row[0]
            if isinstance(cell.value,datetime):
                fmt=re.sub(r'"[^"]*"|\\.',"",cell.number_format).lower()
                # 地区标记中的 zh-CN、en-US 不表示小时或秒。
                elapsed=bool(re.search(r"\[(?:h+|s+)\]",fmt))
                fmt=re.sub(r"\[[^\]]*\]","",fmt)
                result[cell.row-2]=elapsed or bool(re.search(r"h|s",fmt)) or cell.value.time()!=time()
        return result
    finally:
        book.close()


def _row(df, index, row):
    result=row.to_dict()
    result["来源文件"]=df.attrs.get("source_path","")
    result["来源工作表"]=df.attrs.get("source_sheet","")
    result["原始行号"]=int(index)+2 if isinstance(index,int) else str(index)
    return result


def _time_screen(df, options, notes):
    column=options.get("column")
    if column not in df.columns:
        raise ValueError(f"没有所选日期列 {column}")
    if not any(options.get(k) for k in ("night","weekend","holiday")):
        raise ValueError("没有选择夜间、周末或节假日条件")
    start,end=time.fromisoformat(options.get("start","08:00")),time.fromisoformat(options.get("end","18:00"))
    if options.get("night") and start==end:
        raise ValueError("正常工作时间的开始和结束不能相同")
    calendar={}
    if options.get("holiday") or options.get("weekend"):
        try:
            calendar=load_calendar()
        except (OSError,ValueError,KeyError,TypeError) as exc:
            notes.append({"检查":"入账时间","说明":f"日历未能读取，休息日检查未做 {exc}"})
    precision=_precision(df,column) if options.get("night") else {}
    selected=[]
    counts=defaultdict(int)
    for index,row in df.iterrows():
        stamp,has_time=_datetime(row[column])
        if stamp is None:
            counts["日期缺失或不是完整日期，未检查"]+=1
            continue
        has_time=precision.get(index,has_time)
        reasons=[]
        kind,holiday,source=day_type(stamp.date(),calendar) if calendar else ("日历缺失","","")
        if options.get("holiday") or options.get("weekend"):
            if kind=="日历缺失":
                counts[f"{stamp.year}年日历缺失，该年周末调休及节假日未检查"]+=1
            elif options.get("weekend") and kind=="普通周末":
                reasons.append("普通周末")
            elif options.get("holiday") and kind in ("法定节假日","假期连休或调休"):
                reasons.append(holiday+" "+kind)
        if options.get("night"):
            if not has_time:
                counts["未提供真实时分秒，夜间检查未做"]+=1
            else:
                t=stamp.time()
                in_work=start<=t<end if start<end else (t>=start or t<end)
                if not in_work:
                    reasons.append("夜间或指定工作时段外")
        if reasons:
            result=_row(df,index,row)
            result.update({"筛选原因":"；".join(reasons),"所用时间列":column,"日历依据":source})
            selected.append(result)
    notes.extend({"检查":"入账时间","说明":k,"记录数":v} for k,v in counts.items())
    notes.append({"检查":"入账时间","说明":f"日期列 {column}；夜间 {bool(options.get('night'))}；正常时段 {start} 至 {end}；周末 {bool(options.get('weekend'))}；节假日 {bool(options.get('holiday'))}","记录数":len(selected)})
    return pd.DataFrame(selected) if selected else pd.DataFrame({"提示":["没有选中记录，请同时查看筛选说明中的未检查项目"]})


def _split_screen(df, options, notes):
    subject_col=options.get("subject_column")
    subjects=options.get("subjects",[])
    direction=options.get("direction")
    if subject_col not in df or not subjects or direction not in ("借方","贷方"):
        raise ValueError("请选要查的科目列、科目及借方或贷方")
    try:
        limit=Decimal(str(options.get("limit","")).replace(",",""))
        minimum=Decimal(str(options.get("minimum","")).replace(",",""))
        if not limit.is_finite() or not minimum.is_finite() or not 0<minimum<limit:
            raise ValueError()
    except (InvalidOperation,ValueError):
        raise ValueError("审批限额和最低候选金额必须是正数，最低金额须小于限额")
    amount_col=direction+"发生额"
    if amount_col not in df:
        raise ValueError(f"缺少 {amount_col}")
    date_col,party_col=options.get("date_column"),options.get("party_column")
    can_group=date_col in df and party_col in df
    if not can_group:
        notes.append({"检查":"拆分审批","说明":"未选择有效日期列和交易对方列，仅列金额候选"})
    ledger=df.attrs.get("ledger_column")
    if not ledger:
        from src.io.reader import detect_ledger_account_column
        ledger=detect_ledger_account_column(list(df.columns))
    groups=defaultdict(list)
    selected=[]
    red=0
    selected_subjects={str(x).strip() for x in subjects}
    for index,row in df.iterrows():
        if str(row[subject_col]).strip() not in selected_subjects:
            continue
        try:
            amount=Decimal(str(row[amount_col]).replace(",",""))
            if not amount.is_finite():
                raise InvalidOperation()
        except InvalidOperation:
            notes.append({"检查":"拆分审批","说明":f"原始行 {index} 金额无法识别，未检查"})
            continue
        if amount<0:
            red+=1
            continue
        if not minimum<=amount<limit:
            continue
        result=_row(df,index,row)
        result.update({"筛选金额":amount,"审批限额":limit,"候选最低金额":minimum,
                       "同组金额":None,"同组记录数":None,"筛选原因":"金额候选，尚未归集"})
        selected.append(result)
        if can_group:
            stamp,_=_datetime(row[date_col])
            party="" if pd.isna(row[party_col]) else str(row[party_col]).strip()
            ledger_value="" if not ledger else row[ledger]
            if stamp and party and (not ledger or pd.notna(ledger_value) and str(ledger_value).strip()):
                groups[(str(ledger_value),stamp.date(),party)].append(result)
            else:
                notes.append({"检查":"拆分审批","说明":f"原始行 {index} 的日期、对方或账套缺失，仅作金额候选"})
    for number,(key,records) in enumerate(groups.items(),1):
        total=sum((r["筛选金额"] for r in records),Decimal(0))
        for result in records:
            result.update({"归集组号":number,"归集日期":key[1],"归集对方":key[2],
                           "同组金额":total,"同组记录数":len(records),
                           "筛选原因":"多笔候选合计超过限额，待核查" if len(records)>1 and total>limit else "金额候选"})
    if red:
        notes.append({"检查":"拆分审批","说明":"所选方向存在红字，保留于原始数据，不取绝对值或抵减候选金额；核查时应结合冲销记录","记录数":red})
    notes.append({"检查":"拆分审批","说明":f"取 {subject_col} 中的 {'、'.join(map(str,subjects))}，{amount_col}；金额范围 [{minimum}, {limit})。同组金额每行重复展示，不可再次求和。结果仅供核查。","记录数":len(selected)})
    return pd.DataFrame(selected) if selected else pd.DataFrame({"提示":["没有选中记录，请同时查看筛选说明中的未检查项目"]})


def run_screening(df, options=None):
    if not options:
        return {}
    sheets,notes={},[]
    for key,title,handler in (("time","入账时间筛选",_time_screen),("split","拆分审批筛选",_split_screen)):
        if key not in options:
            continue
        try:
            sheets[title]=handler(df,options[key],notes)
        except (ValueError,KeyError,OSError,TypeError) as exc:
            notes.append({"检查":title,"说明":f"本项未检查 {exc}"})
            sheets[title]=pd.DataFrame({"提示":[f"本项未检查 {exc}"]})
    if notes:
        sheets["筛选说明"]=pd.DataFrame(notes)
    return sheets


def add_screening_arguments(parser):
    parser.add_argument("--screen-date",help="入账时间筛选所用的原始日期或时间列")
    parser.add_argument("--screen-weekend",action="store_true",help="筛普通周末，排除调休补班")
    parser.add_argument("--screen-holiday",action="store_true",help="筛法定节假日及假期调休")
    parser.add_argument("--screen-night",action="store_true",help="筛正常工作时段外的记录")
    parser.add_argument("--work-start",default="08:00",help="正常工作开始时间 HH:MM")
    parser.add_argument("--work-end",default="18:00",help="正常工作结束时间 HH:MM")
    parser.add_argument("--split-subject-column",help="拆分筛选所用科目列")
    parser.add_argument("--split-subject",action="append",help="要查的科目原值，可重复指定")
    parser.add_argument("--split-direction",choices=["借方","贷方"],help="只读取这一方向的原始金额")
    parser.add_argument("--approval-limit",help="本次审批限额")
    parser.add_argument("--split-minimum",help="接近限额的最低候选金额")
    parser.add_argument("--split-date",help="归集所用日期列")
    parser.add_argument("--split-party",help="归集所用交易对方列")


def options_from_args(args):
    options={}
    if args.screen_date or args.screen_night or args.screen_weekend or args.screen_holiday:
        options["time"]={"column":args.screen_date,"night":args.screen_night,"weekend":args.screen_weekend,
                         "holiday":args.screen_holiday,"start":args.work_start,"end":args.work_end}
    if args.split_subject_column or args.split_subject or args.approval_limit or args.split_direction or args.split_minimum or args.split_date or args.split_party:
        options["split"]={"subject_column":args.split_subject_column,"subjects":args.split_subject or [],
                          "direction":args.split_direction,"limit":args.approval_limit,"minimum":args.split_minimum,
                          "date_column":args.split_date,"party_column":args.split_party}
    return options
